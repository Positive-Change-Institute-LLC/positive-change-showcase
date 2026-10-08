"""Generate a Genesis Relics launch kit from locally supplied concept art.

This is a one-shot local asset generator, not a background service. It does not
publish content, process payments, create payment QR codes, or connect to AI,
wallet, marketplace, or blockchain services.
"""

import argparse
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


TURNKEY_NAME = "GenesisRelics"
CONCEPT_ART_FILES = ("GenesisRelics_render1.png", "GenesisRelics_render2.png")
NARRATION_SCRIPTS = (
    "Genesis Relics explores proposed NFT evolution, dashboards, and membership tiers.",
    "A concept for multi-wallet tracking and cinematic visuals; no wallet integration is active.",
    "Explore proposed rarity progression without implying investment performance or returns.",
)
PRICING_TABLE = (
    ("Tier", "Proposed price", "Proposed benefits"),
    ("One-Time Purchase", "$6,995", "Turnkey access; scope and availability require confirmation"),
    ("VIP Monthly Membership", "$69/month", "Proposed analytics and support; not an active offer"),
    ("VIP Annual Membership", "$699/year", "Proposed membership benefits; not an active offer"),
    ("In-App Items", "To be determined", "No purchases or transactions are implemented"),
)
BRANDING = "Genesis Relics — Positive Change Institute LLC"
DISCLOSURE = "Concept only • Prices are proposed, not an appraisal or active offer"


def create_promo_graphic(concept_art_file, narration_excerpt, output_file):
    """Create a labeled promo image from an existing local image."""
    with Image.open(concept_art_file) as source:
        image = source.convert("RGBA")
    draw = ImageDraw.Draw(image, "RGBA")
    font = ImageFont.load_default()
    margin = max(12, image.width // 40)
    panel_height = min(max(72, image.height // 4), image.height)
    draw.rectangle(
        (0, 0, image.width, panel_height),
        fill=(0, 0, 0, 190),
    )
    text_width = max(1, image.width - margin * 2)
    lines = []
    for paragraph in (BRANDING, narration_excerpt[:240], DISCLOSURE):
        words = paragraph.split()
        line = ""
        for word in words:
            candidate = f"{line} {word}".strip()
            if draw.textbbox((0, 0), candidate, font=font)[2] <= text_width:
                line = candidate
            elif line:
                lines.append(line)
                line = word
        if line:
            lines.append(line)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 3
    max_lines = max(1, (panel_height - margin) // line_height)
    y = margin
    for line in lines[:max_lines]:
        draw.text((margin, y), line, font=font, fill=(255, 255, 255, 255))
        y += line_height
    output = Path(output_file)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.convert("RGB").save(output, format="PNG")
    return output


def _add_textbox(slide, text, left, top, width, height, font_size, bold=False):
    shape = slide.shapes.add_textbox(left, top, width, height)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(25, 40, 65)
    return shape


def create_investor_deck(output_file, art_files):
    """Create a fresh concept deck so repeated runs never append duplicate slides."""
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank_layout)
    _add_textbox(slide, f"{TURNKEY_NAME} — BlueChip Turnkey", Inches(0.7), Inches(1.5),
                 Inches(8), Inches(0.7), 28, True)
    _add_textbox(slide, "Presented by Positive Change Institute LLC", Inches(0.7), Inches(2.3),
                 Inches(8), Inches(0.5), 18)
    _add_textbox(slide, "Concept launch kit • No live AI, wallet, NFT, or payment integration",
                 Inches(0.7), Inches(3), Inches(8.5), Inches(0.8), 14)

    slide = presentation.slides.add_slide(blank_layout)
    _add_textbox(slide, "Proposed Pricing & Membership Options", Inches(0.5), Inches(0.35),
                 Inches(9), Inches(0.6), 24, True)
    _add_textbox(slide, "Proposed prices only; availability and terms are unverified. Not investment advice.",
                 Inches(0.5), Inches(0.95), Inches(9), Inches(0.45), 11)
    table = slide.shapes.add_table(
        len(PRICING_TABLE), len(PRICING_TABLE[0]),
        Inches(0.5), Inches(1.6), Inches(9), Inches(3.0),
    ).table
    for row_index, row in enumerate(PRICING_TABLE):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            if row_index == 0:
                cell.text_frame.paragraphs[0].font.bold = True

    for index, narration in enumerate(NARRATION_SCRIPTS, start=1):
        slide = presentation.slides.add_slide(blank_layout)
        _add_textbox(slide, f"Concept Narration {index}", Inches(0.7), Inches(0.7),
                     Inches(8), Inches(0.6), 24, True)
        _add_textbox(slide, narration, Inches(0.7), Inches(1.6),
                     Inches(8.5), Inches(1.5), 18)
        _add_textbox(slide, "Illustrative concept copy; capabilities are not represented as live.",
                     Inches(0.7), Inches(3.2), Inches(8.5), Inches(0.6), 12)

    slide = presentation.slides.add_slide(blank_layout)
    _add_textbox(slide, "Cinematic Concept Art", Inches(0.5), Inches(0.35),
                 Inches(9), Inches(0.6), 24, True)
    if art_files:
        for index, image_path in enumerate(art_files):
            with Image.open(image_path) as image:
                image_width, image_height = image.size
            box_width = box_height = 4.0
            scale = min(box_width / image_width, box_height / image_height)
            rendered_width = box_width * scale
            rendered_height = box_height * scale
            slide.shapes.add_picture(
                str(image_path),
                Inches(0.7 + index * 4.5 + (box_width - rendered_width) / 2),
                Inches(1.3 + (box_height - rendered_height) / 2),
                width=Inches(rendered_width),
                height=Inches(rendered_height),
            )
    else:
        _add_textbox(slide, "No concept-art files were found in the selected source directory.",
                     Inches(0.7), Inches(1.5), Inches(8), Inches(0.7), 16)

    output = Path(output_file)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return output


def process_turnkey_queue(source_dir, output_dir):
    """Generate one launch kit, skipping missing art and never running in a loop."""
    source = Path(source_dir).resolve()
    root = Path(output_dir).resolve()
    turnkey_path = root / TURNKEY_NAME
    concept_path = turnkey_path / "ConceptArt"
    social_path = turnkey_path / "SocialGraphics"
    deck_path = turnkey_path / "LaunchDeck" / "BlueChip_Overview.pptx"
    concept_path.mkdir(parents=True, exist_ok=True)
    social_path.mkdir(parents=True, exist_ok=True)

    integrated_art = []
    missing_art = []
    for filename in CONCEPT_ART_FILES:
        source_file = source / filename
        if not source_file.is_file():
            missing_art.append(filename)
            continue
        copied_file = concept_path / filename
        shutil.copy2(source_file, copied_file)
        integrated_art.append(copied_file)

    graphics = []
    for index, image_path in enumerate(integrated_art):
        excerpt = NARRATION_SCRIPTS[index % len(NARRATION_SCRIPTS)]
        graphics.append(create_promo_graphic(
            image_path, excerpt, social_path / f"{TURNKEY_NAME}_promo_{index + 1}.png"
        ))

    create_investor_deck(deck_path, integrated_art)
    return {
        "turnkey_path": turnkey_path,
        "concept_art": integrated_art,
        "missing_art": missing_art,
        "social_graphics": graphics,
        "investor_deck": deck_path,
        "payment_qr_codes": [],
        "background_worker_started": False,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir", default=".",
        help="Directory containing GenesisRelics_render1.png and GenesisRelics_render2.png",
    )
    parser.add_argument(
        "--output-dir", default="BlueChip_Turnkeys_LaunchKit",
        help="Output directory for generated launch-kit files",
    )
    args = parser.parse_args()
    result = process_turnkey_queue(args.source_dir, args.output_dir)
    for label, value in result.items():
        print(f"{label}: {value}")


if __name__ == "__main__":
    main()
