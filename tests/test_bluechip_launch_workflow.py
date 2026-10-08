import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation

from BlueChip_Turnkeys_AutoQueue import process_turnkey_queue


class BlueChipLaunchWorkflowTests(unittest.TestCase):
    def test_missing_art_is_reported_and_deck_is_still_created(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            result = process_turnkey_queue(root / "empty", root / "output")

            self.assertEqual(len(result["missing_art"]), 2)
            self.assertEqual(result["concept_art"], [])
            self.assertEqual(result["social_graphics"], [])
            self.assertFalse(result["background_worker_started"])
            self.assertEqual(result["payment_qr_codes"], [])
            deck = Presentation(str(result["investor_deck"]))
            self.assertEqual(len(deck.slides), 6)
            self.assertTrue(any(
                "No concept-art files" in shape.text
                for shape in deck.slides[-1].shapes
                if shape.has_text_frame
            ))

    def test_supplied_art_creates_graphics_and_repeatable_deck(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source"
            source.mkdir()
            Image.new("RGB", (320, 180), "navy").save(source / "GenesisRelics_render1.png")
            output = root / "output"

            result = process_turnkey_queue(source, output)
            first_slide_count = len(Presentation(str(result["investor_deck"])).slides)
            repeated = process_turnkey_queue(source, output)

            self.assertEqual(len(result["concept_art"]), 1)
            self.assertEqual(len(result["social_graphics"]), 1)
            self.assertTrue(result["social_graphics"][0].is_file())
            self.assertTrue(repeated["investor_deck"].is_file())
            self.assertEqual(
                len(Presentation(str(repeated["investor_deck"])).slides),
                first_slide_count,
            )

    def test_script_does_not_define_payment_links_or_background_loop(self):
        script = Path(__file__).parents[1] / "BlueChip_Turnkeys_AutoQueue.py"
        source = script.read_text(encoding="utf-8")
        self.assertNotIn("PAYMENT_LINKS", source)
        self.assertNotIn("while True", source)


if __name__ == "__main__":
    unittest.main()
